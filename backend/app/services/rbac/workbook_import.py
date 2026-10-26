"""One-time workbook import with a preview, and a guarded reset (ADR 0005 D3).

Reads every column of each managed sheet (hidden ones too, as Excel's CSV
export does) and records the sheet's header text and order in ``rbac_sheets``
so the export reproduces it. The workbook is the source of truth, so rows are
copied as they are and never checked against the access policy; the only
findings are about the copy itself (empty sheets, unrecognised columns,
identical rows). With RBAC_IMPORT_REMOVE_DUPLICATES on, a row identical to an
earlier row on its sheet is left out (and reported); the rows kept are
numbered 0..n-1 in workbook order, so the export is the workbook without them.
"""
from __future__ import annotations

import hashlib
import io
import logging
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from typing import Any, Dict, List, Optional, Tuple

from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.rbac import (
    PLANE_METASTORE,
    RULE_ACTIVE,
    RULE_DELETED,
    SOURCE_IMPORT,
    RbacImportModel,
    RbacRuleModel,
    RbacSheetModel,
)
from app.models.rbac import (
    CANONICAL_FIELDS,
    Finding,
    ImportPreview,
    ImportResult,
    ImportSheetSummary,
    ImportSummary,
    SheetColumn,
)
from app.services.rbac import audit, inflight, layout

logger = logging.getLogger(__name__)

CONFIRMATION_PHRASE = "RESET ACCESS RULES"

MODE_INITIAL = "initial"
MODE_RESET = "reset"
STATUS_PREVIEWED = "previewed"
STATUS_COMPLETED = "completed"
STATUS_FAILED = "failed"

# Findings kept on the import row, and returned with a preview.
MAX_STORED_FINDINGS = 500
MAX_PREVIEW_FINDINGS = 200

UNREADABLE = "That file isn't a readable .xlsx workbook."

_CANONICAL = frozenset(CANONICAL_FIELDS)


class ImportFailed(Exception):
    """The workbook can't be imported, with a message fit for the user."""


@dataclass
class _Row:
    position: int
    cells: Dict[str, str]
    extra: Dict[str, str]
    environment: Optional[str]
    row_hash: str


@dataclass
class _Sheet:
    spec: layout.SheetSpec
    title: str
    columns: List[Dict[str, Optional[str]]]
    unknown_headers: List[str]
    rows: List[_Row] = field(default_factory=list)
    duplicates: int = 0
    duplicates_removed: bool = False
    warnings: int = 0


@dataclass
class _Workbook:
    sheets: List[_Sheet]
    skipped: List[str]
    missing: List[str]
    findings: List[Finding]


def _utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _cell_text(value: Any) -> str:
    """A cell's text as Excel's CSV export writes it for General-formatted cells.

    Number formats aren't applied: a date is written as ISO 8601 (Excel would use
    the cell's format, e.g. ``1/31/2026``) and a percentage or currency cell as
    its plain number. The confirmed sheets hold text only.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        # Excel keeps 15 significant digits.
        return format(value, ".15g").replace("e", "E")
    if isinstance(value, datetime):
        if value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    return str(value)


def _read_sheet(ws: Any, spec: layout.SheetSpec, *, remove_duplicates: bool) -> Tuple[_Sheet, List[Finding]]:
    findings: List[Finding] = []
    width = ws.max_column or 0
    max_row = ws.max_row or 0
    first = next(ws.iter_rows(min_row=1, max_row=1, max_col=width, values_only=True), ()) if width else ()
    headers = [_cell_text(v) for v in first]

    if max_row <= 1 and not any(headers):
        sheet = _Sheet(spec=spec, title=ws.title, columns=layout.default_columns(spec.plane), unknown_headers=[],
                       duplicates_removed=remove_duplicates)
        findings.append(Finding(
            severity="warn", code="empty_sheet", sheet=spec.name,
            message="The sheet is empty. Its export will have the default headers and no rows.",
        ))
        return sheet, findings

    fields = [f if f in _CANONICAL else None for f in layout.map_headers(headers)]
    columns = [{"header": h, "field": f} for h, f in zip(headers, fields)]
    unknown = [h for h, f in zip(headers, fields) if f is None and h.strip()]
    sheet = _Sheet(spec=spec, title=ws.title, columns=columns, unknown_headers=unknown,
                   duplicates_removed=remove_duplicates)
    for header in unknown:
        findings.append(Finding(
            severity="warn", code="unknown_header", sheet=spec.name,
            message=(f'Column "{header}" isn\'t a recognised column (or repeats one). '
                     "Its values are kept as they are and exported unchanged."),
        ))

    # Findings name the row as Excel numbers it: the header is row 1, and blank rows still count.
    first_seen: Dict[str, int] = {}
    for excel_row, values in enumerate(
            ws.iter_rows(min_row=2, max_row=max_row, max_col=width, values_only=True), start=2):
        texts = [_cell_text(v) for v in values]
        # Excel would write an empty row in the used range as ",,,"; it carries no rule.
        if not any(texts):
            continue
        cells = {f: "" for f in CANONICAL_FIELDS}
        extra: Dict[str, str] = {}
        for idx, text in enumerate(texts):
            fld = fields[idx]
            if fld:
                cells[fld] = text
            elif text:
                extra[str(idx)] = text
        environment = spec.environment
        if spec.plane == PLANE_METASTORE:
            environment = layout.resolve_environment(cells["environment_value"])
        row_hash = audit.compute_row_hash(spec.name, cells, extra)
        if row_hash in first_seen:
            sheet.duplicates += 1
            if remove_duplicates:
                findings.append(Finding(
                    severity="warn", code="duplicate_row_removed", sheet=spec.name, row_number=excel_row,
                    message=f"Row {excel_row} is identical to row {first_seen[row_hash]} and was removed.",
                ))
                continue
            findings.append(Finding(
                severity="warn", code="duplicate_row", sheet=spec.name, row_number=excel_row,
                message=(f"Row {excel_row} is identical to row {first_seen[row_hash]}. "
                         "Both are kept so the export matches the workbook."),
            ))
        else:
            first_seen[row_hash] = excel_row
        position = len(sheet.rows)
        sheet.rows.append(_Row(position=position, cells=cells, extra=extra, environment=environment, row_hash=row_hash))
    return sheet, findings


def _parse(data: bytes, *, remove_duplicates: bool) -> _Workbook:
    try:
        import openpyxl
        from openpyxl.worksheet.worksheet import Worksheet

        # Not read-only: hidden columns and the used range need the full model.
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse error means "not a workbook we can read"
        logger.info("Uploaded workbook couldn't be read: %s", exc)
        raise ImportFailed(UNREADABLE) from exc

    try:
        sheets: List[_Sheet] = []
        skipped: List[str] = []
        sheet_findings: Dict[str, List[Finding]] = {}
        seen: set = set()
        for title in wb.sheetnames:
            ws = wb[title]
            spec = layout.sheet_spec(title)
            if spec is None or spec.name in seen or not isinstance(ws, Worksheet):
                skipped.append(title)
                continue
            seen.add(spec.name)
            sheet, found = _read_sheet(ws, spec, remove_duplicates=remove_duplicates)
            sheets.append(sheet)
            sheet_findings[spec.name] = found
    except Exception as exc:  # noqa: BLE001
        logger.warning("Uploaded workbook couldn't be read", exc_info=True)
        raise ImportFailed(UNREADABLE) from exc
    finally:
        wb.close()

    if not sheets:
        names = ", ".join(s.name for s in layout.managed_sheets()[:3])
        raise ImportFailed(f"None of the workbook's tabs is a managed sheet (expected names like {names}).")

    findings: List[Finding] = []
    for sheet in sheets:
        found = sheet_findings[sheet.spec.name]
        sheet.warnings = len(found)
        findings.extend(found)

    missing = [s.name for s in layout.managed_sheets() if s.name not in seen]
    return _Workbook(sheets=sheets, skipped=skipped, missing=missing, findings=findings)


def _summary(sheet: _Sheet) -> ImportSheetSummary:
    return ImportSheetSummary(
        sheet=sheet.spec.name,
        plane=sheet.spec.plane,
        environment=sheet.spec.environment,
        rows=len(sheet.rows),
        columns=[SheetColumn(**c) for c in sheet.columns],
        unknown_headers=sheet.unknown_headers,
        duplicates=sheet.duplicates,
        duplicates_removed=sheet.duplicates_removed,
        warnings=sheet.warnings,
    )


def _remove_duplicates() -> bool:
    return bool(settings.RBAC_IMPORT_REMOVE_DUPLICATES)


def _previewed_with_removal(imp: RbacImportModel) -> bool:
    """Whether the preview removed duplicates (stored per sheet; previews from before the setting kept them)."""
    return any(bool((s or {}).get("duplicates_removed")) for s in (imp.sheets or {}).values())


def _active_rule_count(db: Session) -> int:
    return db.query(RbacRuleModel).filter(RbacRuleModel.status == RULE_ACTIVE).count()


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def preview(db: Session, data: bytes, filename: Optional[str], *, actor: str) -> ImportPreview:
    """Parse the workbook, store an ``rbac_imports`` row (status ``previewed``) and report what an import would do."""
    remove_duplicates = _remove_duplicates()
    parsed = _parse(data, remove_duplicates=remove_duplicates)
    existing = _active_rule_count(db)
    mode = MODE_INITIAL if existing == 0 else MODE_RESET
    summaries = [_summary(s) for s in parsed.sheets]
    imp = RbacImportModel(
        id=audit.new_id(),
        filename=filename,
        file_sha256=_sha256(data),
        mode=mode,
        status=STATUS_PREVIEWED,
        actor=actor,
        sheets={
            s.sheet: {**s.model_dump(mode="json", exclude={"sheet"}), "title": sheet.title}
            for s, sheet in zip(summaries, parsed.sheets)
        },
        skipped_sheets=list(parsed.skipped),
        findings=[f.model_dump(mode="json") for f in parsed.findings[:MAX_STORED_FINDINGS]],
        created_at=_utcnow(),
    )
    db.add(imp)
    in_flight = inflight.count_open(db)
    db.commit()
    logger.info("Previewed workbook import %s (%s): %d sheets, %d rows, %d duplicates (%s), %d findings, mode %s",
                imp.id, filename, len(summaries), sum(s.rows for s in summaries),
                sum(s.duplicates for s in summaries), "removed" if remove_duplicates else "kept",
                len(parsed.findings), mode)
    return ImportPreview(
        import_id=imp.id,
        filename=filename,
        file_sha256=imp.file_sha256,
        sheets=summaries,
        skipped_sheets=parsed.skipped,
        missing_sheets=parsed.missing,
        findings=parsed.findings[:MAX_PREVIEW_FINDINGS],
        findings_total=len(parsed.findings),
        existing_rules=existing,
        mode=mode,
        reset_allowed=bool(settings.RBAC_ALLOW_RESET),
        in_flight_requests=in_flight,
        confirmation_phrase=CONFIRMATION_PHRASE if mode == MODE_RESET else None,
        duplicates_removed=remove_duplicates,
    )


def _check_run(db: Session, imp: Optional[RbacImportModel], data: bytes, mode: str, confirmation: str) -> None:
    if imp is None or imp.status != STATUS_PREVIEWED:
        raise ImportFailed("That preview wasn't found or was already used. Preview the workbook again.")
    if _sha256(data) != imp.file_sha256:
        raise ImportFailed("The uploaded file isn't the one that was previewed. Preview it again before importing.")
    if mode not in (MODE_INITIAL, MODE_RESET):
        raise ImportFailed(f'Unknown import mode "{mode}".')
    if _previewed_with_removal(imp) != _remove_duplicates():
        raise ImportFailed('"Remove duplicate rows on import" was changed after this preview. '
                           "Preview the workbook again.")
    current = MODE_INITIAL if _active_rule_count(db) == 0 else MODE_RESET
    if mode != current:
        if current == MODE_RESET:
            raise ImportFailed("The ledger has rules now, so importing would be a reset. Preview the workbook again.")
        raise ImportFailed("The ledger is empty now, so this is an initial import. Preview the workbook again.")
    if mode == MODE_RESET:
        if not settings.RBAC_ALLOW_RESET:
            raise ImportFailed("Reset is turned off. It can be allowed under Admin -> Settings -> Access Rules.")
        if confirmation != CONFIRMATION_PHRASE:
            raise ImportFailed(f'Type "{CONFIRMATION_PHRASE}" exactly to confirm the reset.')
        open_requests = inflight.count_open(db)
        if open_requests:
            noun = "access request is" if open_requests == 1 else "access requests are"
            raise ImportFailed(f"{open_requests} {noun} still in progress. Wait until they finish, "
                               "or reject them, before resetting.")


def _write(db: Session, imp: RbacImportModel, parsed: _Workbook, mode: str,
           name: str, actor: str) -> Tuple[str, int, int]:
    """Soft-delete (reset) and insert in the caller's transaction. Returns (change set id, created, deleted)."""
    now = _utcnow()
    verb = "reset" if mode == MODE_RESET else "import"
    cs = audit.new_change_set(db, source=SOURCE_IMPORT, source_ref=imp.id, actor=actor,
                              reason=f"Workbook {verb}: {name}")
    touched_sheets: set = set()
    touched_planes: set = set()

    deleted = 0
    if mode == MODE_RESET:
        for rule in db.query(RbacRuleModel).filter(RbacRuleModel.status == RULE_ACTIVE).all():
            before = audit.snapshot(rule)
            rule.status = RULE_DELETED
            rule.version = (rule.version or 1) + 1
            rule.updated_at = now
            rule.updated_by = actor
            audit.record_event(db, cs, rule_id=rule.id, action="delete", before=before, after=audit.snapshot(rule))
            touched_sheets.add(rule.sheet)
            touched_planes.add(rule.plane)
            deleted += 1

    rules: List[RbacRuleModel] = []
    for sheet in parsed.sheets:
        spec = sheet.spec
        for row in sheet.rows:
            rules.append(RbacRuleModel(
                id=audit.new_id(),
                sheet=spec.name,
                plane=spec.plane,
                environment=row.environment,
                position=row.position,
                extra=row.extra,
                status=RULE_ACTIVE,
                version=1,
                row_hash=row.row_hash,
                source=SOURCE_IMPORT,
                source_ref=imp.id,
                created_at=now,
                created_by=actor,
                updated_at=now,
                updated_by=actor,
                **row.cells,
            ))
        touched_sheets.add(spec.name)
        touched_planes.add(spec.plane)
    db.add_all(rules)
    for rule in rules:
        audit.record_event(db, cs, rule_id=rule.id, action="create", before=None, after=audit.snapshot(rule))

    for sheet in parsed.sheets:
        spec = sheet.spec
        recorded = db.get(RbacSheetModel, spec.name)
        if recorded is None:
            recorded = RbacSheetModel(name=spec.name)
            db.add(recorded)
        recorded.plane = spec.plane
        recorded.environment = spec.environment
        recorded.columns = [dict(c) for c in sheet.columns]
        recorded.source_import_id = imp.id
        recorded.updated_at = now

    removed = sum(s.duplicates for s in parsed.sheets if s.duplicates_removed)
    audit.finalize_change_set(cs, created=len(rules), updated=0, deleted=deleted,
                              sheets=list(touched_sheets), planes=list(touched_planes),
                              import_id=imp.id, mode=mode, duplicates_removed=removed)
    imp.status = STATUS_COMPLETED
    imp.mode = mode
    imp.change_set_id = cs.id
    imp.completed_at = now
    imp.error = None
    return cs.id, len(rules), deleted


def _mark_failed(db: Session, import_id: str, error: str) -> None:
    try:
        imp = db.get(RbacImportModel, import_id)
        if imp is None:
            return
        imp.status = STATUS_FAILED
        imp.error = error[:2000]
        imp.completed_at = _utcnow()
        db.commit()
    except Exception:  # noqa: BLE001 - the original failure is what the caller reports
        logger.exception("Couldn't record the failure of workbook import %s", import_id)
        db.rollback()


def run(
    db: Session,
    data: bytes,
    filename: Optional[str],
    *,
    preview_id: str,
    mode: str,
    confirmation: str,
    actor: str,
) -> ImportResult:
    """Import the previewed workbook (the upload's sha256 must match the preview's).

    Duplicates are removed or kept as the preview did; if RBAC_IMPORT_REMOVE_DUPLICATES
    has changed since, the run is refused so the admin previews again.

    ``initial`` requires an empty ledger. ``reset`` requires RBAC_ALLOW_RESET,
    ``confirmation == CONFIRMATION_PHRASE`` and no open rbac_access_change
    request; it soft-deletes every active rule and inserts the workbook's rows
    as one change set (source ``import``). Platform Admin is checked by the caller.
    """
    imp = (
        db.query(RbacImportModel)
        .filter(RbacImportModel.id == preview_id)
        .with_for_update()
        .one_or_none()
    )
    _check_run(db, imp, data, mode, confirmation)
    remove_duplicates = _previewed_with_removal(imp)
    parsed = _parse(data, remove_duplicates=remove_duplicates)
    name = filename or imp.filename or "workbook"

    try:
        change_set_id, created, deleted = _write(db, imp, parsed, mode, name, actor)
        db.commit()
    except Exception as exc:
        db.rollback()
        logger.exception("Workbook import %s failed", preview_id)
        _mark_failed(db, preview_id, f"{exc.__class__.__name__}: {exc}")
        raise ImportFailed("The import failed and nothing was changed. The error is recorded on the import.") from exc

    logger.info("Workbook %s %s by %s: %d rows created, %d deleted, %d duplicates %s (change set %s)",
                mode, preview_id, actor, created, deleted, sum(s.duplicates for s in parsed.sheets),
                "removed" if remove_duplicates else "kept", change_set_id)
    try:
        from app.services.rbac.csv_export import request_export_soon

        request_export_soon()
    except Exception:  # noqa: BLE001 - the import is committed; the schedule still exports it
        logger.warning("Couldn't request an export after workbook import %s", preview_id, exc_info=True)

    return ImportResult(
        import_id=preview_id,
        change_set_id=change_set_id,
        created=created,
        deleted=deleted,
        sheets=[_summary(s) for s in parsed.sheets],
        duplicates_removed=remove_duplicates,
    )


def list_imports(db: Session, *, limit: int = 20) -> List[ImportSummary]:
    rows = (
        db.query(RbacImportModel)
        .order_by(RbacImportModel.created_at.desc())
        .limit(limit)
        .all()
    )
    return [
        ImportSummary(
            id=r.id,
            filename=r.filename,
            mode=r.mode,
            status=r.status,
            actor=r.actor,
            created_at=r.created_at,
            completed_at=r.completed_at,
            change_set_id=r.change_set_id,
            error=r.error,
        )
        for r in rows
    ]
