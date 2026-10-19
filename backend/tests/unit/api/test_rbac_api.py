"""The Access Rules API (``/api/v1/rbac``): feature gate, rights per type, error
shapes, the two-person rule, uploads and exports (ADR 0005 D2, D3, D8, D12, D13).

Each test gets a private in-memory database: the ledger commits and rolls back
its own transactions, which would discard the shared ``db_session`` fixture's
outer transaction. Group membership is mocked through ``User.entitlements``.
"""
import copy
import io
import tempfile

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import app.db  # noqa: F401  (registers every model on Base.metadata)
from app.api.deps import get_current_user
from app.api.v1 import rbac as rbac_api
from app.core.config import _yaml_config, settings
from app.core.default_config import DEFAULT_CONFIG
from app.core.exceptions import PermanentError
from app.db import RequestModel
from app.db.base import Base
from app.db.rbac import RbacChangeSetModel, RbacExportModel, RbacRuleModel
from app.db.session import get_db
from app.models.rbac import ChangeSetIn, CompareResult, ExportRun, ImportPreview, PathCheckResult
from app.models.user import User
from app.services.rbac import approvers, audit, csv_export, layout, ledger, workbook_import
from app.services.rbac import requests as rq
from app.services.rbac.csv_export import ExportBusy
from app.state_machines.facts import add_fact

UC_DEV = "UC_Securables_RBAC_Dev"
UC_PRD = "UC_Securables_RBAC_Prd"
WS_DEV = "WS_Securables_RBAC_Dev"
META = "UC_metastore_RBAC"
OWNER = "metastore_admins"
REASON = "Quarterly access review clean-up"
UC_GROUP = "UC Approvers"
WS_GROUP = "WS Approvers"
MS_GROUP = "Metastore Approvers"
BASE = "/api/v1/rbac"


def _user(email, roles=(), groups=()):
    return User(id=email, email=email, full_name=email, roles=list(roles),
                entitlements=[email, *groups], is_active=True)


ADMIN = _user("admin@example.com", roles=["Platform Admin"])
VIEWER = _user("viewer@example.com", roles=["Governance Admin"])
UC_MEMBER = _user("uc@example.com", groups=["uc approvers"])
WS_MEMBER = _user("ws@example.com", groups=[WS_GROUP])
EVERY_GROUP = _user("every@example.com", groups=[UC_GROUP, WS_GROUP, MS_GROUP])
OUTSIDER = _user("outsider@example.com", groups=["Other Group"])

ENVIRONMENTS = [
    {"key": "dev", "aliases": "development", "sheet_suffix": "Dev", "export_path": "", "self_service": True,
     "write_privileges": True, "service_principals": True, "two_person_rule": False},
    {"key": "prd", "aliases": "prod,production", "sheet_suffix": "Prd", "export_path": "", "self_service": True,
     "write_privileges": False, "service_principals": False, "two_person_rule": False},
]


@pytest.fixture(autouse=True)
def _settings(monkeypatch):
    monkeypatch.setitem(_yaml_config.setdefault("features", {}), "rbac", True)
    monkeypatch.setattr(settings, "RBAC_ENVIRONMENTS", [dict(e) for e in ENVIRONMENTS])
    monkeypatch.setattr(settings, "RBAC_UC_SHEET_PATTERN", "UC_Securables_RBAC_{suffix}")
    monkeypatch.setattr(settings, "RBAC_WS_SHEET_PATTERN", "WS_Securables_RBAC_{suffix}")
    monkeypatch.setattr(settings, "RBAC_METASTORE_SHEET", META)
    monkeypatch.setattr(settings, "RBAC_METASTORE_EXPORT_PATH", "")
    monkeypatch.setattr(settings, "RBAC_UC_APPROVER_GROUP", UC_GROUP)
    monkeypatch.setattr(settings, "RBAC_WS_APPROVER_GROUP", WS_GROUP)
    monkeypatch.setattr(settings, "RBAC_METASTORE_APPROVER_GROUP", MS_GROUP)
    monkeypatch.setattr(settings, "RBAC_VIEWER_ROLES", ["Governance Admin"])
    monkeypatch.setattr(settings, "RBAC_ALLOW_RESET", False)
    monkeypatch.setattr(settings, "RBAC_MAX_CHANGES_PER_COMMIT", 200)
    monkeypatch.setattr(settings, "RBAC_MIN_REASON_LENGTH", 10)
    monkeypatch.setattr(settings, "RBAC_MIN_JUSTIFICATION_LENGTH", 20)
    monkeypatch.setattr(settings, "RBAC_ALLOWED_PRIVILEGES",
                        copy.deepcopy(DEFAULT_CONFIG["rbac"]["allowed_privileges"]))
    monkeypatch.setattr(settings, "RBAC_SYSTEM_GROUPS", ["account users", "users", "admins"])
    monkeypatch.setattr(settings, "RBAC_NAMING_RULES", [])
    monkeypatch.setattr(settings, "RBAC_HEADER_ALIASES", [])
    monkeypatch.setattr(settings, "RBAC_DEFAULT_OWNER", "")
    monkeypatch.setattr(settings, "RBAC_EXPORT_ON_CHANGE", False)
    monkeypatch.setattr(settings, "RBAC_CSV_BOM", False)
    monkeypatch.setattr(settings, "RBAC_CSV_LINE_ENDING", "crlf")
    monkeypatch.setattr(settings, "RBAC_EXPORT_ARCHIVE_DAYS", 90)
    monkeypatch.setattr(csv_export, "_change_pending_since", None)

    async def group_exists(name):
        return True

    monkeypatch.setattr(approvers, "group_exists", group_exists)


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


USERS = {u.email: u for u in (ADMIN, VIEWER, UC_MEMBER, WS_MEMBER, EVERY_GROUP, OUTSIDER)}


@pytest.fixture
def api(db):
    """``api(user)``: a client signed in as ``user``."""
    app = FastAPI()
    app.include_router(rbac_api.router, prefix=BASE)

    def current_user(request: Request) -> User:
        return USERS[request.headers["x-test-user"]]

    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = current_user
    clients = {}

    def client_for(user):
        if user.email not in clients:
            clients[user.email] = TestClient(app, headers={"x-test-user": user.email})
        return clients[user.email]

    yield client_for
    for client in clients.values():
        client.close()


def seed(db, sheet=UC_DEV, *, position=1, principals="g_analysts", privileges="USE_SCHEMA;SELECT",
         kind="schema", catalog="sales_dev", schema="orders", obj="", workspace="", owner=OWNER):
    spec = layout.sheet_spec(sheet)
    cells = dict(
        environment_value=spec.environment or "", securable_type=kind, catalog=catalog, schema_name=schema,
        object_name=obj, workspace_name=workspace, owner=owner, principal_type="group",
        principal_name=principals, privilege=privileges,
    )
    rule = RbacRuleModel(
        id=audit.new_id(), sheet=sheet, plane=spec.plane, environment=spec.environment, position=position,
        extra={}, status="active", version=1, row_hash=audit.compute_row_hash(sheet, cells, {}),
        source="import", **cells,
    )
    db.add(rule)
    db.commit()
    return rule


def seed_ws(db, *, position=1):
    return seed(db, WS_DEV, position=position, kind="workspace", catalog="", schema="", obj="ws-dev",
                workspace="ws-dev", owner="", principals="g_ws", privileges="USER")


def uc_create(**cells):
    row = dict(environment_value="dev", securable_type="schema", catalog="sales_dev", schema_name="payments",
               owner=OWNER, principal_type="group", principal_name="g_new", privilege="USE_SCHEMA;SELECT")
    row.update(cells)
    return {"op": "create", "sheet": UC_DEV, "cells": row}


def ws_create():
    return {"op": "create", "sheet": WS_DEV, "cells": dict(
        environment_value="dev", principal_type="group", principal_name="g_ws_new", workspace_name="ws-dev",
        securable_type="workspace", object_name="ws-dev", privilege="USER")}


def body(*changes, reason=REASON):
    return {"changes": list(changes), "reason": reason}


def rule_count(db):
    db.expire_all()
    return db.query(RbacRuleModel).count()


def xlsx(sheets=None) -> bytes:
    """A small workbook with one UC sheet and a scratch tab."""
    wb = Workbook()
    wb.remove(wb.active)
    sheets = sheets or {
        UC_DEV: [
            ["Environment", "Securable Type", "Catalog", "SchemaName", "ObjectName", "Owner",
             "PrincipalType", "PrincipalName", "AccessPrivilege"],
            ["dev", "schema", "cat_dev", "sales", None, OWNER, "group", "g_one;g_two", "USE_SCHEMA;SELECT"],
            ["dev", "view", "cat_dev", "sales", "daily", OWNER, "group", "g_three", "SELECT"],
        ],
        "Sheet1": [["scratch"]],
    }
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(name, data, content_type="application/octet-stream"):
    return {"file": (name, data, content_type)}


# --- Feature gate ------------------------------------------------------------------------


def test_everything_is_404_when_the_feature_is_off(api, monkeypatch):
    monkeypatch.setitem(_yaml_config["features"], "rbac", False)
    client = api(ADMIN)
    assert client.get(f"{BASE}/permissions").status_code == 404
    assert client.get(f"{BASE}/rules").status_code == 404
    assert client.post(f"{BASE}/changes", json=body(uc_create())).status_code == 404
    assert client.post(f"{BASE}/imports/preview", files=upload("rbac.xlsx", xlsx())).status_code == 404
    assert client.get(f"{BASE}/exports/e/files/f").status_code == 404


# --- Permissions ----------------------------------------------------------------------------


@pytest.mark.parametrize("user, view, imp, export, editable", [
    (ADMIN, True, True, True, {"uc", "uc_metastore", "workspace"}),
    (VIEWER, True, False, False, set()),
    (UC_MEMBER, True, False, True, {"uc"}),
    (OUTSIDER, False, False, False, set()),
])
def test_permissions_per_kind_of_user(api, user, view, imp, export, editable):
    res = api(user).get(f"{BASE}/permissions")
    assert res.status_code == 200
    perms = res.json()
    assert (perms["can_view"], perms["can_import"], perms["can_export"], perms["can_reset"]) == (
        view, imp, export, False)
    assert {p["plane"] for p in perms["planes"] if p["can_edit"]} == editable
    by_plane = {p["plane"]: p for p in perms["planes"]}
    assert by_plane["uc"]["approver_group"] == UC_GROUP
    assert by_plane["uc"]["group_status"] == "found"


@pytest.mark.parametrize("method, path, kwargs", [
    ("get", "/sheets", {}),
    ("get", "/rules", {}),
    ("post", "/changes/validate", {"json": body(uc_create())}),
    ("post", "/changes", {"json": body(uc_create())}),
    ("get", "/history", {}),
    ("get", "/history/cs-1", {}),
    ("post", "/history/cs-1/revert/validate", {"json": {"reason": REASON}}),
    ("post", "/history/cs-1/revert", {"json": {"reason": REASON}}),
    ("get", "/imports", {}),
    ("post", "/imports/preview", {"files": upload("rbac.xlsx", b"x")}),
    ("post", "/imports", {"files": upload("rbac.xlsx", b"x"), "data": {"preview_id": "p", "mode": "initial"}}),
    ("get", "/exports", {}),
    ("post", "/exports", {"json": {"force": True}}),
    ("get", "/exports/e-1/files/f-1", {}),
    ("post", "/exports/test-path", {"json": {"target": "dev"}}),
    ("post", "/exports/compare", {"files": upload("a.csv", b"x"), "data": {"sheet": UC_DEV}}),
])
def test_an_outsider_gets_403_with_a_plain_message_everywhere_else(api, method, path, kwargs):
    res = getattr(api(OUTSIDER), method)(f"{BASE}{path}", **kwargs)
    assert res.status_code == 403
    assert isinstance(res.json()["detail"], str)


# --- Sheets and rules ---------------------------------------------------------------------


def test_sheets_mark_the_members_planes_editable(api, db):
    seed(db, position=1)
    sheets = {s["name"]: s for s in api(UC_MEMBER).get(f"{BASE}/sheets").json()}
    assert list(sheets) == [UC_DEV, UC_PRD, WS_DEV, "WS_Securables_RBAC_Prd", META]
    assert sheets[UC_DEV]["editable"] is True and sheets[UC_DEV]["row_count"] == 1
    assert sheets[WS_DEV]["editable"] is False
    assert sheets[META]["export_target"] == "metastore"
    assert not any(s["editable"] for s in api(VIEWER).get(f"{BASE}/sheets").json())


def test_sheets_carry_the_dropdown_options(api, monkeypatch):
    envs = [dict(e) for e in ENVIRONMENTS]
    envs[0].update(workspaces="ws-dev, ws-dev-analytics", catalogs="sales_dev, finance_dev")
    monkeypatch.setattr(settings, "RBAC_ENVIRONMENTS", envs)
    sheets = {s["name"]: s["options"] for s in api(VIEWER).get(f"{BASE}/sheets").json()}

    uc = sheets[UC_DEV]
    assert set(uc) == {"environments", "securable_types", "principal_types", "privileges", "workspaces", "catalogs"}
    assert uc["environments"] == ["dev"]
    assert uc["securable_types"][:3] == ["catalog", "schema", "table"]
    assert "connection" not in uc["securable_types"] and set(uc["privileges"]) == set(uc["securable_types"])
    assert "SELECT" in uc["privileges"]["schema"] and "MANAGE" not in uc["privileges"]["schema"]
    assert uc["principal_types"] == {"dev": ["group", "service principal"]}
    assert (uc["catalogs"], uc["workspaces"]) == (["sales_dev", "finance_dev"], [])

    ws = sheets[WS_DEV]
    assert "workspace entitlement" in ws["securable_types"]
    assert ws["privileges"]["workspace"] == ["ADMIN", "USER"]
    assert (ws["workspaces"], ws["catalogs"]) == (["ws-dev", "ws-dev-analytics"], [])
    assert sheets["WS_Securables_RBAC_Prd"]["principal_types"] == {"prd": ["group"]}
    assert sheets["WS_Securables_RBAC_Prd"]["workspaces"] == []

    meta = sheets[META]
    assert meta["environments"] == ["dev", "prd"]
    assert meta["principal_types"] == {"dev": ["group", "service principal"], "prd": ["group"]}
    assert "connection" in meta["securable_types"] and (meta["workspaces"], meta["catalogs"]) == ([], [])


def test_a_privilege_added_to_the_allow_list_is_offered_and_accepted(api, monkeypatch):
    client = api(UC_MEMBER)
    change = body(uc_create(privilege="USE_SCHEMA;SELECT;NEW_PRIVILEGE"))
    preview = client.post(f"{BASE}/changes/validate", json=change).json()
    assert preview["blocked"] is True
    assert [f["privilege"] for f in preview["findings"] if f["code"] == "privilege_not_allowed"] == ["NEW_PRIVILEGE"]

    rows = copy.deepcopy(DEFAULT_CONFIG["rbac"]["allowed_privileges"])
    schema = next(r for r in rows if r["securable_type"] == "schema")
    schema["privileges"] += ", NEW_PRIVILEGE"
    monkeypatch.setattr(settings, "RBAC_ALLOWED_PRIVILEGES", rows)
    assert client.post(f"{BASE}/changes/validate", json=change).json()["blocked"] is False
    options = {s["name"]: s["options"] for s in client.get(f"{BASE}/sheets").json()}[UC_DEV]
    assert options["privileges"]["schema"][-1] == "NEW_PRIVILEGE"


def test_a_workspace_off_the_configured_list_warns_on_a_direct_edit(api, db, monkeypatch):
    envs = [dict(e) for e in ENVIRONMENTS]
    envs[0]["workspaces"] = "ws-dev-analytics"
    monkeypatch.setattr(settings, "RBAC_ENVIRONMENTS", envs)
    preview = api(WS_MEMBER).post(f"{BASE}/changes/validate", json=body(ws_create())).json()
    flagged = [f for f in preview["findings"] if f["code"] == "value_not_in_options"]
    assert {(f["field"], f["severity"]) for f in flagged} == {("workspace_name", "warn"), ("object_name", "warn")}
    assert preview["blocked"] is False and "ws-dev-analytics" in flagged[0]["message"]


def test_rules_filter_search_and_cap_the_page_size(api, db):
    seed(db, position=1)
    seed(db, position=2, schema="returns", principals="g_finance")
    seed_ws(db)
    client = api(VIEWER)

    page = client.get(f"{BASE}/rules", params={"sheet": UC_DEV}).json()
    assert page["total"] == 2
    assert [r["schema_name"] for r in page["rules"]] == ["orders", "returns"]
    assert page["rules"][0]["pending_request_ids"] == []
    assert client.get(f"{BASE}/rules", params={"q": "finance"}).json()["total"] == 1
    assert client.get(f"{BASE}/rules", params={"plane": "workspace"}).json()["total"] == 1
    assert client.get(f"{BASE}/rules", params={"environment": "development"}).json()["total"] == 3
    assert client.get(f"{BASE}/rules", params={"limit": 1, "offset": 1}).json()["total"] == 3
    assert client.get(f"{BASE}/rules", params={"limit": 5000}).status_code == 200
    assert client.get(f"{BASE}/rules", params={"limit": 5001}).status_code == 422
    assert client.get(f"{BASE}/rules", params={"plane": "nope"}).status_code == 422


# --- Direct edits ------------------------------------------------------------------------------


def test_a_uc_member_commits_uc_changes_but_not_workspace_ones(api, db):
    client = api(UC_MEMBER)

    preview = client.post(f"{BASE}/changes/validate", json=body(uc_create()))
    assert preview.status_code == 200
    assert preview.json()["blocked"] is False and preview.json()["diffs"][0]["op"] == "create"
    assert rule_count(db) == 0

    res = client.post(f"{BASE}/changes", json=body(uc_create()))
    assert res.status_code == 200
    result = res.json()
    assert result["change_set_id"] and result["request_id"] is None
    assert rule_count(db) == 1
    created = db.get(RbacRuleModel, result["diffs"][0]["rule_id"])
    assert (created.sheet, created.created_by, created.source) == (UC_DEV, UC_MEMBER.email, "admin_edit")

    for path in ("/changes/validate", "/changes"):
        denied = client.post(f"{BASE}{path}", json=body(ws_create()))
        assert denied.status_code == 403
        assert "Workspace securables" in denied.json()["detail"]
        mixed = client.post(f"{BASE}{path}", json=body(uc_create(schema_name="other"), ws_create()))
        assert mixed.status_code == 403
    assert rule_count(db) == 1
    assert api(VIEWER).post(f"{BASE}/changes", json=body(uc_create())).status_code == 403


def test_a_ws_member_updates_a_workspace_row(api, db):
    row = seed_ws(db)
    res = api(WS_MEMBER).post(f"{BASE}/changes", json=body(
        {"op": "update", "rule_id": row.id, "version": 1, "cells": {"principal_name": "g_ws;g_more"}}))
    assert res.status_code == 200
    db.expire_all()
    assert db.get(RbacRuleModel, row.id).principal_name == "g_ws;g_more"
    denied = api(UC_MEMBER).post(f"{BASE}/changes", json=body(
        {"op": "delete", "rule_id": row.id, "version": 2}))
    assert denied.status_code == 403


def test_commit_needs_a_reason_and_returns_the_preview_on_422(api, db):
    client = api(ADMIN)
    # The dry run runs before a reason is written, so a missing reason doesn't block it.
    preview = client.post(f"{BASE}/changes/validate", json=body(uc_create(), reason=""))
    assert preview.status_code == 200 and preview.json()["blocked"] is False
    assert "reason_too_short" not in {f["code"] for f in preview.json()["findings"]}
    short = client.post(f"{BASE}/changes/validate", json=body(uc_create(), reason="tidy"))
    assert short.json()["blocked"] is True

    res = client.post(f"{BASE}/changes", json=body(uc_create(), reason=""))
    assert res.status_code == 422
    res = client.post(f"{BASE}/changes", json=body(uc_create(), reason="tidy"))
    assert res.status_code == 422
    detail = res.json()["detail"]
    assert isinstance(detail["message"], str)
    assert detail["preview"]["blocked"] is True
    assert "reason_too_short" in {f["code"] for f in detail["preview"]["findings"]}
    assert len(detail["preview"]["diffs"]) == 1
    assert rule_count(db) == 0


def test_a_stale_version_is_409_with_the_current_row(api, db):
    row = seed(db)
    client = api(ADMIN)
    first = client.post(f"{BASE}/changes", json=body(
        {"op": "update", "rule_id": row.id, "version": 1, "cells": {"principal_name": "g_x"}}))
    assert first.status_code == 200

    stale = {"op": "update", "rule_id": row.id, "version": 1, "cells": {"owner": "someone_else"}}
    for path in ("/changes/validate", "/changes"):
        res = client.post(f"{BASE}{path}", json=body(stale))
        assert res.status_code == 409
        detail = res.json()["detail"]
        assert detail["rule_id"] == row.id
        assert isinstance(detail["message"], str)
        assert detail["current"]["version"] == 2 and detail["current"]["principal_name"] == "g_x"
        assert detail["current"]["id"] == row.id
    db.expire_all()
    assert db.get(RbacRuleModel, row.id).owner == OWNER


def test_ledger_errors_are_400_with_a_plain_message(api):
    client = api(ADMIN)
    unknown = client.post(f"{BASE}/changes", json=body({"op": "delete", "rule_id": "missing", "version": 1}))
    assert unknown.status_code == 400 and "doesn't exist" in unknown.json()["detail"]
    sheet = client.post(f"{BASE}/changes/validate", json=body({"op": "create", "sheet": "Sheet1", "cells": {}}))
    assert sheet.status_code == 400 and isinstance(sheet.json()["detail"], str)
    assert client.post(f"{BASE}/changes", json={"changes": [], "reason": REASON}).status_code == 422


# --- Two-person rule -------------------------------------------------------------------------


@pytest.fixture
def two_person_dev(monkeypatch):
    envs = [dict(e) for e in ENVIRONMENTS]
    envs[0]["two_person_rule"] = True
    monkeypatch.setattr(settings, "RBAC_ENVIRONMENTS", envs)


@pytest.fixture
def filed(monkeypatch):
    calls = []

    def submit_row_changes(db, change, *, plane, requester_email):
        calls.append({"change": change, "plane": plane, "requester_email": requester_email})
        return f"req-{len(calls)}"

    monkeypatch.setattr("app.services.rbac.requests.submit_row_changes", submit_row_changes)
    return calls


def test_two_person_rule_files_a_request_and_leaves_the_ledger_untouched(api, db, two_person_dev, filed):
    row = seed(db)
    preview = api(UC_MEMBER).post(f"{BASE}/changes/validate", json=body(uc_create())).json()
    assert preview["two_person_required"] is True

    res = api(UC_MEMBER).post(f"{BASE}/changes", json=body(
        uc_create(), {"op": "update", "rule_id": row.id, "version": 1, "cells": {"principal_name": "g_y"}}))
    assert res.status_code == 200
    result = res.json()
    assert (result["request_id"], result["change_set_id"]) == ("req-1", None)
    assert [d["op"] for d in result["diffs"]] == ["create", "update"]

    [call] = filed
    assert (call["plane"], call["requester_email"]) == ("uc", UC_MEMBER.email)
    assert call["change"].reason == REASON and len(call["change"].changes) == 2
    assert rule_count(db) == 1
    db.expire_all()
    assert db.get(RbacRuleModel, row.id).principal_name == "g_analysts"
    assert db.query(RbacChangeSetModel).count() == 0


def test_two_person_rule_still_blocks_and_needs_one_type(api, db, two_person_dev, filed):
    client = api(ADMIN)
    blocked = client.post(f"{BASE}/changes", json=body(uc_create(), reason="tidy"))
    assert blocked.status_code == 422 and blocked.json()["detail"]["preview"]["blocked"] is True

    mixed = client.post(f"{BASE}/changes", json=body(uc_create(), ws_create()))
    assert mixed.status_code == 400
    assert mixed.json()["detail"].startswith("Split the change by type")
    assert filed == [] and rule_count(db) == 0


def test_environments_without_the_rule_commit_directly(api, db, two_person_dev, filed):
    res = api(ADMIN).post(f"{BASE}/changes", json=body(
        {"op": "create", "sheet": UC_PRD, "cells": uc_create(environment_value="prd", catalog="sales_prd")["cells"]}))
    assert res.status_code == 200 and res.json()["change_set_id"] and res.json()["request_id"] is None
    assert filed == [] and rule_count(db) == 1


# --- History and revert ---------------------------------------------------------------------


def test_history_detail_and_revert(api, db):
    client = api(UC_MEMBER)
    committed = client.post(f"{BASE}/changes", json=body(uc_create())).json()
    cs_id = committed["change_set_id"]
    created_id = committed["diffs"][0]["rule_id"]

    page = api(VIEWER).get(f"{BASE}/history").json()
    assert page["total"] == 1
    assert page["change_sets"][0]["id"] == cs_id and page["change_sets"][0]["summary"]["planes"] == ["uc"]
    detail = api(VIEWER).get(f"{BASE}/history/{cs_id}").json()
    assert detail["actor"] == UC_MEMBER.email and [d["op"] for d in detail["diffs"]] == ["create"]
    assert api(VIEWER).get(f"{BASE}/history/missing").status_code == 404
    assert api(VIEWER).post(f"{BASE}/history/{cs_id}/revert", json={"reason": REASON}).status_code == 403
    assert api(WS_MEMBER).post(f"{BASE}/history/{cs_id}/revert/validate", json={}).status_code == 403

    preview = client.post(f"{BASE}/history/{cs_id}/revert/validate", json={"reason": ""})
    assert preview.status_code == 200
    assert [(d["op"], d["rule_id"]) for d in preview.json()["diffs"]] == [("delete", created_id)]

    short = client.post(f"{BASE}/history/{cs_id}/revert", json={"reason": "oops"})
    assert short.status_code == 422 and short.json()["detail"]["preview"]["blocked"] is True

    reverted = client.post(f"{BASE}/history/{cs_id}/revert", json={"reason": "Created on the wrong schema"})
    assert reverted.status_code == 200
    revert_id = reverted.json()["change_set_id"]
    db.expire_all()
    assert db.get(RbacRuleModel, created_id).status == "deleted"
    assert db.get(RbacChangeSetModel, cs_id).reverted_by == revert_id
    assert api(VIEWER).get(f"{BASE}/history").json()["total"] == 2

    again = client.post(f"{BASE}/history/{cs_id}/revert", json={"reason": "Created on the wrong schema"})
    assert again.status_code == 400 and "already reverted" in again.json()["detail"]
    assert client.post(f"{BASE}/history/missing/revert", json={"reason": REASON}).status_code == 404


def test_revert_needs_edit_rights_on_every_type_it_touches(api, db):
    mixed = api(ADMIN).post(f"{BASE}/changes", json=body(uc_create(), ws_create())).json()
    cs_id = mixed["change_set_id"]
    for user in (UC_MEMBER, WS_MEMBER):
        assert api(user).post(f"{BASE}/history/{cs_id}/revert/validate", json={}).status_code == 403
    assert api(EVERY_GROUP).post(f"{BASE}/history/{cs_id}/revert/validate", json={}).status_code == 200


def test_a_revert_after_a_later_edit_is_409(api, db):
    row = seed(db)
    client = api(ADMIN)
    first = client.post(f"{BASE}/changes", json=body(
        {"op": "update", "rule_id": row.id, "version": 1, "cells": {"principal_name": "g_x"}})).json()
    client.post(f"{BASE}/changes", json=body(
        {"op": "update", "rule_id": row.id, "version": 2, "cells": {"owner": "new_owner"}}))
    res = client.post(f"{BASE}/history/{first['change_set_id']}/revert/validate", json={})
    assert res.status_code == 409
    assert res.json()["detail"]["rule_id"] == row.id and res.json()["detail"]["current"]["version"] == 3


@pytest.fixture
def steps(db, monkeypatch):
    """Run the workflow's plan and apply steps against the test database."""
    monkeypatch.setattr(db, "close", lambda: None)
    monkeypatch.setattr(rq, "_session", lambda: db)


def _original(db, *changes):
    """An earlier edit to revert; commit_change_set itself ignores the two-person rule."""
    return ledger.commit_change_set(db, ChangeSetIn.model_validate(body(*changes)), actor=ADMIN.email)


def _approve(db, request_id, user):
    add_fact(db, request_id, "approval_received", {"approved_by": user.email}, actor=user.email)


def test_two_person_revert_files_a_request_and_leaves_the_ledger_untouched(api, db, two_person_dev):
    original = _original(db, uc_create())
    cs_id, created_id = original.change_set_id, original.diffs[0].rule_id
    client = api(UC_MEMBER)

    short = client.post(f"{BASE}/history/{cs_id}/revert", json={"reason": "oops"})
    assert short.status_code == 422 and short.json()["detail"]["preview"]["blocked"] is True

    res = client.post(f"{BASE}/history/{cs_id}/revert", json={"reason": "Created on the wrong schema"})
    assert res.status_code == 200
    result = res.json()
    assert result["change_set_id"] is None and result["request_id"]
    assert [(d["op"], d["rule_id"]) for d in result["diffs"]] == [("delete", created_id)]
    db.expire_all()
    assert db.get(RbacRuleModel, created_id).status == "active"
    assert db.get(RbacChangeSetModel, cs_id).reverted_by is None
    assert db.query(RbacChangeSetModel).count() == 1
    request = db.get(RequestModel, result["request_id"])
    assert request.requester_email == UC_MEMBER.email
    assert (request.state_context["revert_of"], request.state_context["plane"]) == (cs_id, "uc")
    [rule] = api(VIEWER).get(f"{BASE}/rules", params={"sheet": UC_DEV}).json()["rules"]
    assert rule["pending_request_ids"] == [result["request_id"]]

    again = client.post(f"{BASE}/history/{cs_id}/revert", json={"reason": "Created on the wrong schema"})
    assert again.status_code == 400 and "already open" in again.json()["detail"]
    assert db.query(RequestModel).count() == 1


def test_a_filed_revert_is_recorded_once_another_member_approves(api, db, two_person_dev, steps):
    original = _original(db, uc_create())
    cs_id, created_id = original.change_set_id, original.diffs[0].rule_id
    request_id = api(UC_MEMBER).post(f"{BASE}/history/{cs_id}/revert",
                                     json={"reason": "Created on the wrong schema"}).json()["request_id"]

    assert rq.prepare_step(request_id)["approvers"] == [UC_GROUP]
    _approve(db, request_id, UC_MEMBER)
    with pytest.raises(PermanentError, match="someone other than the requester"):
        rq.record_step(request_id)
    db.expire_all()
    assert db.get(RbacRuleModel, created_id).status == "active"

    _approve(db, request_id, EVERY_GROUP)
    recorded = rq.record_step(request_id)
    db.expire_all()
    assert db.get(RbacRuleModel, created_id).status == "deleted"
    assert db.get(RbacChangeSetModel, cs_id).reverted_by == recorded["change_set_id"]
    assert rq.record_step(request_id)["already_recorded"]
    history = api(VIEWER).get(f"{BASE}/history").json()
    assert history["total"] == 2
    revert = next(cs for cs in history["change_sets"] if cs["id"] == recorded["change_set_id"])
    assert (revert["source"], revert["actor"]) == ("revert", UC_MEMBER.email)
    assert revert["reason"] == f"Created on the wrong schema (request {request_id})"


def test_a_two_person_revert_across_types_is_one_request_for_platform_admins(api, db, two_person_dev, steps):
    original = _original(db, uc_create(), ws_create())
    res = api(EVERY_GROUP).post(f"{BASE}/history/{original.change_set_id}/revert", json={"reason": REASON})
    assert res.status_code == 200 and res.json()["request_id"]
    request = db.get(RequestModel, res.json()["request_id"])
    assert request.state_context["plane"] is None
    assert rq.prepare_step(request.id)["approvers"] == []
    assert rule_count(db) == 2 and db.get(RbacChangeSetModel, original.change_set_id).reverted_by is None


# --- Import ------------------------------------------------------------------------------------


def test_import_is_platform_admin_only(api, monkeypatch):
    def fail(*a, **kw):
        raise AssertionError("the service must not be reached")

    monkeypatch.setattr(workbook_import, "preview", fail)
    monkeypatch.setattr(workbook_import, "run", fail)
    for user in (UC_MEMBER, EVERY_GROUP, VIEWER):
        client = api(user)
        assert client.get(f"{BASE}/imports").status_code == 403
        res = client.post(f"{BASE}/imports/preview", files=upload("rbac.xlsx", xlsx()))
        assert res.status_code == 403 and "Platform Admins" in res.json()["detail"]
        assert client.post(f"{BASE}/imports", files=upload("rbac.xlsx", xlsx()),
                           data={"preview_id": "p", "mode": "initial"}).status_code == 403


def test_import_preview_and_run_end_to_end(api, db):
    client = api(ADMIN)
    data = xlsx()
    preview = client.post(f"{BASE}/imports/preview", files=upload("C:\\Users\\me\\rbac.xlsx", data))
    assert preview.status_code == 200
    p = preview.json()
    assert (p["mode"], p["filename"]) == ("initial", "rbac.xlsx")
    assert p["skipped_sheets"] == ["Sheet1"]
    assert {s["sheet"]: s["rows"] for s in p["sheets"]} == {UC_DEV: 2}

    run = client.post(f"{BASE}/imports", files=upload("rbac.xlsx", data),
                      data={"preview_id": p["import_id"], "mode": "initial", "confirmation": ""})
    assert run.status_code == 200
    assert run.json()["created"] == 2 and run.json()["change_set_id"]
    assert rule_count(db) == 2

    again = client.post(f"{BASE}/imports", files=upload("rbac.xlsx", data),
                        data={"preview_id": p["import_id"], "mode": "initial"})
    assert again.status_code == 400 and "already used" in again.json()["detail"]
    missing = client.post(f"{BASE}/imports", files=upload("rbac.xlsx", data), data={"mode": "initial"})
    assert missing.status_code == 400

    imports = client.get(f"{BASE}/imports").json()
    assert [(i["id"], i["status"], i["actor"]) for i in imports] == [(p["import_id"], "completed", ADMIN.email)]


def test_import_rejects_unreadable_files_and_other_extensions(api, monkeypatch):
    client = api(ADMIN)
    csv_upload = client.post(f"{BASE}/imports/preview", files=upload("rbac.csv", b"a,b\r\n"))
    assert csv_upload.status_code == 400 and ".xlsx" in csv_upload.json()["detail"]
    garbage = client.post(f"{BASE}/imports/preview", files=upload("rbac.xlsx", b"not a workbook"))
    assert garbage.status_code == 400 and garbage.json()["detail"] == workbook_import.UNREADABLE
    no_file = client.post(f"{BASE}/imports/preview", files={"other": ("rbac.xlsx", xlsx())})
    assert no_file.status_code == 400 and '"file"' in no_file.json()["detail"]
    two_files = client.post(f"{BASE}/imports/preview", files=[("file", ("a.xlsx", b"a")), ("file", ("b.xlsx", b"b"))])
    assert two_files.status_code == 400
    malformed = client.post(f"{BASE}/imports/preview", content=b"--x\r\nnonsense",
                            headers={"content-type": "multipart/form-data; boundary=x"})
    assert malformed.status_code == 400 and isinstance(malformed.json()["detail"], str)
    not_multipart = client.post(f"{BASE}/imports/preview", content=b"{}", headers={"content-type": "application/json"})
    assert not_multipart.status_code == 400


def test_import_uploads_over_the_limit_are_413(api, monkeypatch):
    def fail(*a, **kw):
        raise AssertionError("the service must not be reached")

    monkeypatch.setattr(workbook_import, "preview", fail)
    monkeypatch.setattr(rbac_api, "MAX_WORKBOOK_BYTES", 2 * 1024 * 1024)
    client = api(ADMIN)
    # Over the file limit but within the form's allowance: caught after reading.
    just_over = client.post(f"{BASE}/imports/preview", files=upload("big.xlsx", b"x" * (2 * 1024 * 1024 + 10)))
    assert just_over.status_code == 413
    assert just_over.json()["detail"] == "The file is larger than the 2 MB limit."
    # Far over: refused from Content-Length before the body is read.
    far_over = client.post(f"{BASE}/imports/preview", files=upload("big.xlsx", b"x" * (3 * 1024 * 1024)))
    assert far_over.status_code == 413


def test_uploads_stay_in_memory(api, monkeypatch):
    """Starlette spools uploads over 1 MB to a temporary file; these must never roll over to disk."""
    def no_disk(self):
        raise AssertionError("upload rolled over to disk")

    monkeypatch.setattr(tempfile.SpooledTemporaryFile, "rollover", no_disk)
    seen = {}

    def preview(db, data, filename, *, actor):
        seen.update(size=len(data), filename=filename, actor=actor)
        return ImportPreview(import_id="imp-1", filename=filename, file_sha256="0" * 64, sheets=[], mode="initial")

    monkeypatch.setattr(workbook_import, "preview", preview)
    payload = b"PK" + b"x" * (5 * 1024 * 1024)
    res = api(ADMIN).post(f"{BASE}/imports/preview", files=upload("rbac.xlsx", payload))
    assert res.status_code == 200
    assert seen == {"size": len(payload), "filename": "rbac.xlsx", "actor": ADMIN.email}


def test_a_platform_admin_can_revert_an_import_but_an_approver_cannot(api, db):
    client = api(ADMIN)
    data = xlsx()
    p = client.post(f"{BASE}/imports/preview", files=upload("rbac.xlsx", data)).json()
    cs_id = client.post(f"{BASE}/imports", files=upload("rbac.xlsx", data),
                        data={"preview_id": p["import_id"], "mode": "initial"}).json()["change_set_id"]

    denied = api(EVERY_GROUP).post(f"{BASE}/history/{cs_id}/revert/validate", json={})
    assert denied.status_code == 403 and "workbook import" in denied.json()["detail"]
    allowed = api(ADMIN).post(f"{BASE}/history/{cs_id}/revert/validate", json={})
    assert allowed.status_code == 200 and len(allowed.json()["diffs"]) == 2


# --- Export ------------------------------------------------------------------------------------


def _run(**kw):
    fields = dict(id="exp-1", trigger="manual", actor=ADMIN.email, status="completed", watermark=3)
    fields.update(kw)
    return ExportRun(**fields)


def test_export_run_needs_an_editable_type_and_maps_busy_to_409(api, monkeypatch):
    calls = []

    def run_export(db, *, trigger, actor=None, force=False):
        calls.append((trigger, actor, force))
        return _run(actor=actor)

    monkeypatch.setattr(csv_export, "run_export", run_export)
    assert api(VIEWER).post(f"{BASE}/exports", json={"force": True}).status_code == 403
    res = api(UC_MEMBER).post(f"{BASE}/exports", json={"force": True})
    assert res.status_code == 200 and res.json()["id"] == "exp-1"
    assert api(ADMIN).post(f"{BASE}/exports").status_code == 200
    assert calls == [("manual", UC_MEMBER.email, True), ("manual", ADMIN.email, False)]

    def busy(db, **kw):
        raise ExportBusy("An export is already running.")

    monkeypatch.setattr(csv_export, "run_export", busy)
    res = api(ADMIN).post(f"{BASE}/exports", json={"force": False})
    assert res.status_code == 409 and res.json()["detail"] == "An export is already running."


def test_export_status_is_for_viewers(api):
    status = api(VIEWER).get(f"{BASE}/exports")
    assert status.status_code == 200
    assert [t["key"] for t in status.json()["targets"]] == ["dev", "prd", "metastore"]
    assert status.json()["runs"] == []


def test_test_path_accepts_only_known_targets(api, monkeypatch):
    calls = []

    def test_path(target):
        calls.append(target)
        return PathCheckResult(target=target, path="/Volumes/x", ok=True, message="ok")

    monkeypatch.setattr(csv_export, "test_path", test_path)
    client = api(WS_MEMBER)
    assert client.post(f"{BASE}/exports/test-path", json={"target": "dev"}).json()["ok"] is True
    assert client.post(f"{BASE}/exports/test-path", json={"target": "metastore"}).status_code == 200
    unknown = client.post(f"{BASE}/exports/test-path", json={"target": "/tmp/anywhere"})
    assert unknown.status_code == 400 and "isn't an export target" in unknown.json()["detail"]
    assert api(VIEWER).post(f"{BASE}/exports/test-path", json={"target": "dev"}).status_code == 403
    assert calls == ["dev", "metastore"]


def test_compare_checks_the_sheet_and_the_size(api, monkeypatch):
    calls = []

    def compare(db, sheet, uploaded):
        calls.append((sheet, uploaded))
        return CompareResult(sheet=sheet, identical=True, ours_sha256="a", theirs_sha256="a")

    monkeypatch.setattr(csv_export, "compare", compare)
    client = api(VIEWER)
    res = client.post(f"{BASE}/exports/compare", files=upload("manual.csv", b"a,b\r\n"),
                      data={"sheet": UC_DEV.lower()})
    assert res.status_code == 200 and res.json()["identical"] is True
    assert calls == [(UC_DEV, b"a,b\r\n")]

    unknown = client.post(f"{BASE}/exports/compare", files=upload("manual.csv", b"a"), data={"sheet": "Sheet1"})
    assert unknown.status_code == 400 and "managed sheets" in unknown.json()["detail"]

    monkeypatch.setattr(rbac_api, "MAX_COMPARE_BYTES", 1024 * 1024)
    big = client.post(f"{BASE}/exports/compare", files=upload("manual.csv", b"x" * (1024 * 1024 + 1)),
                      data={"sheet": UC_DEV})
    assert big.status_code == 413 and big.json()["detail"] == "The file is larger than the 1 MB limit."
    assert len(calls) == 1


def test_download_of_an_unknown_file_is_404(api):
    res = api(VIEWER).get(f"{BASE}/exports/nope/files/nope")
    assert res.status_code == 404 and isinstance(res.json()["detail"], str)


def test_export_to_a_local_folder_and_download_the_file_by_id(api, db, monkeypatch, tmp_path):
    monkeypatch.setattr("app.core.config.dev_features_allowed", lambda: True)
    envs = [dict(e) for e in ENVIRONMENTS]
    envs[0]["export_path"] = str(tmp_path / "dev")
    monkeypatch.setattr(settings, "RBAC_ENVIRONMENTS", envs)
    monkeypatch.setattr(settings, "RBAC_METASTORE_EXPORT_PATH", str(tmp_path / "metastore"))
    client = api(ADMIN)
    client.post(f"{BASE}/changes", json=body(uc_create(), uc_create(
        schema_name="notes", principal_name='g_quoted, "special"')))

    res = client.post(f"{BASE}/exports", json={"force": False})
    assert res.status_code == 200
    run = res.json()
    assert run["status"] == "completed" and run["actor"] == ADMIN.email
    files = {f["sheet"]: f for f in run["files"]}
    assert files["UC_Securables_RBAC_Prd"]["status"] == "skipped"
    uc = files[UC_DEV]
    assert uc["status"] == "written" and uc["downloadable"] is True and uc["row_count"] == 2

    download = api(VIEWER).get(f"{BASE}/exports/{run['id']}/files/{uc['id']}")
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("text/csv")
    assert download.headers["content-disposition"] == f'attachment; filename="{UC_DEV}.csv"'
    assert download.content == (tmp_path / "dev" / f"{UC_DEV}.csv").read_bytes()
    lines = download.content.decode("utf-8").split("\r\n")
    assert lines[0] == "Environment,Securable Type,Catalog,SchemaName,ObjectName,Owner,PrincipalType,PrincipalName,AccessPrivilege"
    assert '"g_quoted, ""special"""' in lines[2]

    assert api(VIEWER).get(f"{BASE}/exports/{run['id']}/files/{files['UC_Securables_RBAC_Prd']['id']}").status_code == 404
    assert api(VIEWER).get(f"{BASE}/exports/other-run/files/{uc['id']}").status_code == 404
    assert api(OUTSIDER).get(f"{BASE}/exports/{run['id']}/files/{uc['id']}").status_code == 403
    assert db.query(RbacExportModel).count() == 1
    status = api(VIEWER).get(f"{BASE}/exports").json()
    assert status["runs"][0]["id"] == run["id"] and status["pending_changes"] is False
