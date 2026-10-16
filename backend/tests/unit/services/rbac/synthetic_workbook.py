"""A synthetic RBAC workbook following the confirmed layouts (ADR 0005, Verification).

Generic names only. Each sheet is a cell matrix (header row first; ``None`` is an
empty cell) plus the columns to hide, so tests can derive what Excel's CSV export
would write for a sheet without going through the importer.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Dict, List, Optional, Sequence, Tuple, Union

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

Cell = Union[str, int, float, bool, None]

UC_HEADERS: List[Cell] = ["Environment", "Securable Type", "Catalog", "SchemaName", "ObjectName", "Owner",
                          "PrincipalType", "PrincipalName", "AccessPrivilege"]
WS_HEADERS: List[Cell] = ["Environment", "PrincipalType", "PrincipalName", "WorkspaceName", "SecurableType",
                          "SecurableName", "SecurablePermission"]
METASTORE_HEADERS: List[Cell] = ["Securable Type", "Catalog", "SchemaName", "Owner", "PrincipalType",
                                 "PrincipalName", "AccessPrivilege"]

OWNER = "grp_metastore_admins"
QUOTED_NOTE = 'Approved by "data office", ticket 12'
MULTILINE_NOTE = "first line\nsecond line"

# (environment key, sheet suffix) as in the default RBAC_ENVIRONMENTS.
ENVIRONMENTS: List[Tuple[str, str]] = [("dev", "Dev"), ("tst", "Tst"), ("stg", "Stg"), ("prd", "Prd")]
METASTORE_SHEET = "UC_metastore_RBAC"


def uc_sheet(suffix: str) -> str:
    return f"UC_Securables_RBAC_{suffix}"


def ws_sheet(suffix: str) -> str:
    return f"WS_Securables_RBAC_{suffix}"


@dataclass
class SyntheticSheet:
    rows: List[List[Cell]]
    # 0-based column indexes to hide.
    hidden: Tuple[int, ...] = ()

    @property
    def width(self) -> int:
        return max(len(r) for r in self.rows)

    def padded(self) -> List[List[Cell]]:
        return [list(r) + [None] * (self.width - len(r)) for r in self.rows]


def _uc_rows(env: str) -> List[List[Cell]]:
    analysts = ";".join(f"grp_{env}_analysts_{i:02d}" for i in range(1, 13))
    return [
        [env, "schema", f"cat_{env}", "sales", None, OWNER, "group", analysts, "USE_SCHEMA;SELECT"],
        [env, "schema", f"cat_{env}", "finance", None, OWNER, "group",
         f"grp_{env}_finance;grp_{env}_audit", "USE_SCHEMA;SELECT;MODIFY"],
        [env, "view", f"cat_{env}", "finance", "fiscal_month", OWNER, "group", f"grp_{env}_reporting", "SELECT"],
        # An sp_ name under PrincipalType "group", as seen in the real sheets.
        [env, "view", f"cat_{env}", "sales", "daily_orders", OWNER, "group", f"sp_{env}_etl", "MODIFY"],
    ]


def _ws_rows(env: str) -> List[List[Cell]]:
    return [
        [env, "group", f"grp_{env}_engineers;grp_{env}_analysts", f"ws_{env}", "workspace", f"ws_{env}", "USER"],
        [env, "group", f"grp_{env}_ws_admins", f"ws_{env}", "workspace", f"ws_{env}", "ADMIN"],
        [env, "group", f"grp_{env}_analysts", f"ws_{env}", "workspace entitlement", f"ws_{env}",
         "workspace_access;databricks_sql_access"],
        [env, "service principal", f"sp_{env}_jobs", f"ws_{env}", "workspace", f"ws_{env}", "USER"],
    ]


def _metastore_rows(with_environment: bool) -> SyntheticSheet:
    rows: List[List[Cell]] = [
        ["stg", "connection", "conn_crm_stg", None, OWNER, "service principal", "sp_stg_ingest",
         "MANAGE;USE_CONNECTION"],
        ["prd", "connection", "conn_crm_prd", None, OWNER, "service principal", "sp_prd_ingest",
         "MANAGE;USE_CONNECTION"],
        ["Dev", "external location", "loc_landing_dev", None, OWNER, "group",
         "grp_dev_engineers;grp_dev_analysts", "READ_FILES"],
        ["production", "storage credential", "cred_lake_prd", None, OWNER, "group", "grp_prd_platform",
         "CREATE_EXTERNAL_LOCATION"],
        ["sandbox", "connection", "conn_scratch", None, OWNER, "group", "grp_dev_engineers", "USE_CONNECTION"],
    ]
    if with_environment:
        return SyntheticSheet(rows=[["Environment", *METASTORE_HEADERS], *rows], hidden=(0,))
    return SyntheticSheet(rows=[list(METASTORE_HEADERS), *[r[1:] for r in rows]])


def synthetic_sheets(*, metastore_environment_column: bool = True) -> Dict[str, SyntheticSheet]:
    """The nine managed sheets (in the workbook's tab order) plus the ``Sheet1`` scratch tab."""
    sheets: Dict[str, SyntheticSheet] = {}
    for env, suffix in ENVIRONMENTS:
        rows = [list(WS_HEADERS), *_ws_rows(env)]
        hidden: Tuple[int, ...] = ()
        if env == "tst":
            # A hidden, unknown column in the middle, holding numbers.
            rows = [r[:3] + [v] + r[3:] for r, v in zip(rows, ["Ticket", 1042, 7.0, None, 1043])]
            hidden = (3,)
        if env == "prd":
            # A stray value under an empty header, past the last named column.
            rows[2] = rows[2] + ["checked by hand"]
        sheets[ws_sheet(suffix)] = SyntheticSheet(rows=rows, hidden=hidden)
    sheets[METASTORE_SHEET] = _metastore_rows(metastore_environment_column)
    for env, suffix in ENVIRONMENTS:
        rows = [list(UC_HEADERS), *_uc_rows(env)]
        hidden = ()
        if env == "dev":
            # A hidden "Notes" column with a value that needs quoting and one with a line break.
            notes = ["Notes", QUOTED_NOTE, None, MULTILINE_NOTE, None]
            rows = [r + [n] for r, n in zip(rows, notes)]
            hidden = (9,)
        if env == "stg":
            rows.insert(2, list(rows[1]))  # a duplicate row
        sheets[uc_sheet(suffix)] = SyntheticSheet(rows=rows, hidden=hidden)
    sheets["Sheet1"] = SyntheticSheet(rows=[["scratch", "ideas"], ["try this", 3]])
    return sheets


def write_workbook(sheets: Dict[str, SyntheticSheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, sheet in sheets.items():
        ws = wb.create_sheet(title)
        for r, row in enumerate(sheet.rows, start=1):
            for c, value in enumerate(row, start=1):
                if value is not None:
                    ws.cell(row=r, column=c, value=value)
        for idx in sheet.hidden:
            ws.column_dimensions[get_column_letter(idx + 1)].hidden = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_workbook(
    *,
    metastore_environment_column: bool = True,
    only: Optional[Sequence[str]] = None,
) -> Tuple[bytes, Dict[str, SyntheticSheet]]:
    """The workbook bytes and the sheets written (``only`` keeps a subset of tabs)."""
    sheets = synthetic_sheets(metastore_environment_column=metastore_environment_column)
    if only is not None:
        sheets = {k: v for k, v in sheets.items() if k in only}
    return write_workbook(sheets), sheets


def managed(sheets: Dict[str, SyntheticSheet]) -> Dict[str, SyntheticSheet]:
    return {k: v for k, v in sheets.items() if k != "Sheet1"}


def data_rows(sheets: Dict[str, SyntheticSheet]) -> int:
    return sum(len(s.rows) - 1 for s in managed(sheets).values())


def import_initial(db, data: bytes, *, filename: str = "rbac.xlsx", actor: str = "admin@example.com"):
    """Preview and run an initial import; returns (preview, result)."""
    from app.services.rbac import workbook_import

    preview = workbook_import.preview(db, data, filename, actor=actor)
    result = workbook_import.run(db, data, filename, preview_id=preview.import_id, mode="initial",
                                 confirmation="", actor=actor)
    return preview, result
